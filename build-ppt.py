from pathlib import Path
import os
import tempfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
OUT = ROOT / os.environ.get("PPT_OUT", "金融客户定制App产品方案-演示版.pptx")
SHOT = ROOT / "screenshots"
ASSET = ROOT / "assets"
TEXT_DIR = Path(tempfile.mkdtemp(prefix="finance-ppt-text-"))
TEXT_FONT = "/System/Library/Fonts/STHeiti Medium.ttc"
TEXT_SCALE = 2

NAVY = RGBColor(7, 27, 67)
BLUE = RGBColor(28, 83, 174)
GOLD = RGBColor(215, 170, 75)
INK = RGBColor(23, 35, 60)
MUTED = RGBColor(104, 116, 139)
PALE = RGBColor(242, 245, 249)
LINE = RGBColor(224, 231, 241)
WHITE = RGBColor(255, 255, 255)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial Unicode MS"):
    """Add native text by default; raster mode is used only for the visual PDF fallback."""
    if os.environ.get("RASTER_TEXT") == "1":
        return raster_text(slide, value, x, y, w, h, size, color, bold, align)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for index, line in enumerate(str(value).split("\n")):
        if index:
            p = tf.add_paragraph()
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                rpr.append(node)
            node.set("typeface", font)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def raster_text(slide, value, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    px_w = max(2, int(w * 144))
    px_h = max(2, int(h * 144))
    font_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    font_size = max(8, int(size * TEXT_SCALE))
    pil_font = ImageFont.truetype(font_path, font_size)
    rgba = (color[0], color[1], color[2], 255)
    canvas = Image.new("RGBA", (px_w, px_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(canvas)
    lines = []
    for paragraph in str(value).split("\n"):
        line = ""
        for char in paragraph:
            candidate = line + char
            if line and draw.textlength(candidate, font=pil_font) > px_w:
                lines.append(line)
                line = char
            else:
                line = candidate
        lines.append(line)
    line_height = max(font_size + 3, pil_font.getbbox("中")[3] + 4)
    top = 0
    for line in lines:
        line_width = draw.textlength(line, font=pil_font)
        left = (px_w - line_width) / 2 if align == PP_ALIGN.CENTER else px_w - line_width if align == PP_ALIGN.RIGHT else 0
        draw.text((left, top), line, font=pil_font, fill=rgba)
        top += line_height
    path = TEXT_DIR / f"raster-text-{len(list(TEXT_DIR.iterdir())):04d}.png"
    canvas.save(path)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def rect(slide, x, y, w, h, color=WHITE, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, color)
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.7)
    return shape


def title(slide, kicker, heading, page):
    text(slide, "金融行业自主私域 App 方案", 0.7, 0.48, 4, 0.22, 9, MUTED, False)
    text(slide, heading, 0.7, 0.82, 8.5, 0.55, 26, NAVY, True)
    text(slide, f"{page:02d}", 12.25, 0.55, 0.45, 0.25, 10, MUTED, True, PP_ALIGN.RIGHT)


def add_phone(slide, filename, x, y, w=2.55, h=5.35):
    """Place the real demo screenshot inside a box without changing its aspect ratio."""
    source = SHOT / filename
    with Image.open(source) as image:
        ratio = image.height / image.width
    shown_h = min(h, w * ratio)
    shown_w = shown_h / ratio
    slide.shapes.add_picture(str(source), Inches(x), Inches(y), width=Inches(shown_w), height=Inches(shown_h))
    return shown_w, shown_h


def add_asset(slide, filename, x, y, w, h):
    """Place an image without distortion, matching the image treatment used in the H5 demo."""
    source = ASSET / filename
    with Image.open(source) as image:
        ratio = image.height / image.width
    shown_h = min(h, w * ratio)
    shown_w = shown_h / ratio
    slide.shapes.add_picture(str(source), Inches(x), Inches(y), width=Inches(shown_w), height=Inches(shown_h))
    return shown_w, shown_h


def add_cover_device(slide, x, y, w=3.6, h=5.8):
    source = SHOT / "01-home-platform.png"
    crop_path = TEXT_DIR / "cover-device.png"
    with Image.open(source) as image:
        image.crop((770, 28, 1260, 925)).save(crop_path)
    with Image.open(crop_path) as image:
        ratio = image.height / image.width
    shown_h = min(h, w * ratio)
    shown_w = shown_h / ratio
    slide.shapes.add_picture(str(crop_path), Inches(x), Inches(y), width=Inches(shown_w), height=Inches(shown_h))


def make_icon(kind, color):
    path = TEXT_DIR / f"icon-{kind}.png"
    image = Image.new("RGBA", (180, 180), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    stroke = (color[0], color[1], color[2], 255)
    width = 8
    if kind == "content":
        draw.rounded_rectangle((42, 25, 138, 155), radius=10, outline=stroke, width=width)
        draw.line((62, 62, 118, 62), fill=stroke, width=width)
        draw.line((62, 91, 118, 91), fill=stroke, width=width)
        draw.line((62, 120, 105, 120), fill=stroke, width=width)
    elif kind == "interaction":
        draw.rounded_rectangle((28, 35, 145, 118), radius=20, outline=stroke, width=width)
        draw.polygon([(58, 118), (48, 148), (83, 118)], outline=stroke, fill=None)
        draw.ellipse((58, 72, 68, 82), fill=stroke)
        draw.ellipse((84, 72, 94, 82), fill=stroke)
        draw.ellipse((110, 72, 120, 82), fill=stroke)
    elif kind == "topic":
        draw.ellipse((30, 30, 150, 150), outline=stroke, width=width)
        draw.ellipse((62, 62, 118, 118), outline=stroke, width=width)
        draw.line((90, 12, 90, 46), fill=stroke, width=width)
        draw.line((90, 134, 90, 168), fill=stroke, width=width)
        draw.line((12, 90, 46, 90), fill=stroke, width=width)
        draw.line((134, 90, 168, 90), fill=stroke, width=width)
    else:
        draw.line((24, 90, 132, 90), fill=stroke, width=width)
        draw.polygon([(126, 58), (160, 90), (126, 122)], outline=stroke, fill=None)
        draw.ellipse((28, 62, 58, 92), outline=stroke, width=width)
        draw.ellipse((58, 88, 88, 118), outline=stroke, width=width)
    image.save(path)
    return path


def add_cover_gradient(slide):
    path = TEXT_DIR / "cover-gradient.png"
    width, height = 1600, 900
    image = Image.new("RGB", (width, height))
    pixels = image.load()
    start = (23, 71, 220)
    end = (29, 151, 238)
    for x in range(width):
        ratio = x / (width - 1)
        color = tuple(int(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3))
        for y in range(height):
            pixels[x, y] = color
    image.save(path)
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))


def add_cover_art(slide):
    path = TEXT_DIR / "cover-art.png"
    width, height = 1600, 900
    image = Image.new("RGB", (width, height))
    pixels = image.load()
    start = (8, 25, 72)
    end = (20, 86, 168)
    for x in range(width):
        ratio = x / (width - 1)
        color = tuple(int(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3))
        for y in range(height):
            pixels[x, y] = color
    draw = ImageDraw.Draw(image, "RGBA")
    # Low-contrast radial arcs keep the cover dimensional without competing with the title.
    for box, alpha in [((780, -240, 1740, 720), 42), ((900, -120, 1860, 840), 30), ((1040, 0, 2000, 960), 22)]:
        draw.ellipse(box, outline=(112, 184, 255, alpha), width=3)
    for x in range(1120, 1600, 80):
        draw.line((x, 0, x - 330, 900), fill=(151, 207, 255, 18), width=2)
    draw.line((80, 775, 850, 775), fill=(224, 180, 83, 230), width=4)
    image.save(path)
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))


def add_content_background(slide):
    """Create a restrained presentation background with a cool paper tone and fine structure lines."""
    path = TEXT_DIR / "content-background.png"
    width, height = 1600, 900
    image = Image.new("RGB", (width, height))
    pixels = image.load()
    left = (250, 251, 253)
    right = (240, 246, 255)
    for x in range(width):
        ratio = x / (width - 1)
        color = tuple(int(left[i] * (1 - ratio) + right[i] * ratio) for i in range(3))
        for y in range(height):
            pixels[x, y] = color
    draw = ImageDraw.Draw(image)
    for offset in (0, 70, 140, 210):
        draw.line((1180 + offset, 0, 1600, 420 + offset), fill=(221, 233, 249), width=2)
    draw.line((0, 870, 1600, 870), fill=(220, 229, 241), width=2)
    draw.line((0, 872, 410, 872), fill=(215, 170, 75), width=3)
    image.save(path)
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))


def bullet(slide, label, body, x, y, w, accent=BLUE):
    rect(slide, x, y + 0.04, 0.07, 0.07, accent, True)
    text(slide, label, x + 0.2, y, w - 0.2, 0.23, 13, NAVY, True)
    text(slide, body, x + 0.2, y + 0.28, w - 0.2, 0.5, 10, MUTED)


def new_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, bg)
    if bg != NAVY:
        add_content_background(slide)
    footer_color = RGBColor(184, 201, 229) if bg == NAVY else RGBColor(153, 163, 178)
    text(slide, "小鹅通·金融BU", 0.7, 7.13, 2.0, 0.18, 8, footer_color, False)
    return slide


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 1. Cover
slide = new_slide(prs, NAVY)
add_cover_art(slide)
text(slide, "金融行业自主私域 App 方案", 0.78, 0.9, 4.8, 0.3, 12, GOLD, True)
text(slide, "金融行业自主私域\nApp 产品方案", 0.78, 1.72, 7.0, 1.35, 36, WHITE, True)
text(slide, "面向金融客户的内容服务、私域运营与品牌化阵地建设", 0.82, 3.58, 7.4, 0.55, 16, RGBColor(210, 225, 247))
text(slide, "圈子运营  ·  消息触达  ·  投教学习  ·  直播服务  ·  会员承接", 0.82, 5.55, 8.5, 0.3, 12, RGBColor(210, 225, 247))
text(slide, "产品方案稿  ·  2026", 0.82, 6.65, 4, 0.25, 10, RGBColor(176, 198, 227))
text(slide, "小鹅通·金融BU", 0.7, 7.13, 2.0, 0.18, 8, RGBColor(210, 226, 250), False)

# 2. Platform dependence
slide = new_slide(prs, PALE)
title(slide, "01 / PLATFORM DEPENDENCE", "金融客户当前的增长，仍高度依赖外部平台", 2)
text(slide, "平台带来流量、内容分发和交易入口，但用户关系并不天然归客户所有。", 0.72, 1.45, 8.8, 0.35, 15, NAVY, True)
for i, (a, b, c) in enumerate([
    ("获客", "短视频 / 直播 / 搜索", BLUE), ("分发", "内容推荐与平台流量", GOLD),
    ("转化", "课程 / 会员 / 服务", RGBColor(38,151,125)), ("关系", "群聊 / 私聊 / 关注", RGBColor(190,78,91)),
]):
    x = 0.85 + i * 3.05
    rect(slide, x, 2.35, 2.55, 1.2, WHITE, True, LINE)
    rect(slide, x, 2.35, 2.55, 0.12, c, True)
    text(slide, a, x + 0.18, 2.62, 2.1, 0.25, 15, NAVY, True)
    text(slide, b, x + 0.18, 3.02, 2.1, 0.3, 11, MUTED)
text(slide, "平台可以带来用户，但客户需要自己的用户经营阵地。", 0.85, 4.45, 7.6, 0.4, 18, BLUE, True)
text(slide, "核心问题不是有没有流量，而是用户关系能否持续沉淀。", 0.85, 5.45, 7.6, 0.32, 14, MUTED)

# 3. Risks
slide = new_slide(prs)
title(slide, "02 / BUSINESS RISK", "平台限制带来的经营风险", 3)
for i, (a, b) in enumerate([
    ("触达风险", "群发、私聊、外链和营销内容存在规则约束"),
    ("内容风险", "金融直播、投顾和收益相关内容审核更严格"),
    ("关系风险", "用户标签、圈子互动和客户归属难以统一沉淀"),
    ("连续性风险", "账号、店铺或直播间异常会导致用户失联"),
]):
    x = 0.85 + (i % 2) * 3.65
    y = 1.75 + (i // 2) * 1.65
    rect(slide, x, y, 3.15, 1.18, WHITE, True, LINE)
    text(slide, a, x + 0.2, y + 0.2, 2.7, 0.25, 15, NAVY, True)
    text(slide, b, x + 0.2, y + 0.57, 2.7, 0.4, 10, MUTED)
rect(slide, 8.45, 1.75, 3.3, 4.45, NAVY, True)
text(slide, "平台规则变化", 8.85, 2.15, 2.5, 0.28, 17, WHITE, True, PP_ALIGN.CENTER)
text(slide, "↓", 9.95, 2.85, 0.3, 0.3, 20, GOLD, True, PP_ALIGN.CENTER)
text(slide, "运营动作受限", 8.9, 3.35, 2.4, 0.28, 17, WHITE, True, PP_ALIGN.CENTER)
text(slide, "↓", 9.95, 4.05, 0.3, 0.3, 20, GOLD, True, PP_ALIGN.CENTER)
text(slide, "用户关系不稳定", 8.8, 4.55, 2.6, 0.28, 17, WHITE, True, PP_ALIGN.CENTER)
text(slide, "风险不是没有流量，而是流量无法沉淀。", 0.85, 5.75, 6.8, 0.35, 16, BLUE, True)

# 4. Impact
slide = new_slide(prs, PALE)
title(slide, "03 / IMPACT", "平台限制如何影响用户、内容和收入", 4)
headers = [("用户", "找不到、触达不到、无法持续学习", BLUE), ("内容", "直播回放、课程资料和观点难沉淀", GOLD), ("收入", "转化链路中断，复购和召回成本上升", RGBColor(190,78,91))]
for i, (a, b, c) in enumerate(headers):
    x = 0.85 + i * 4.05
    rect(slide, x, 1.75, 3.45, 1.65, WHITE, True, LINE)
    rect(slide, x, 1.75, 3.45, 0.12, c, True)
    text(slide, a, x + 0.22, 2.1, 2.8, 0.28, 18, NAVY, True)
    text(slide, b, x + 0.22, 2.6, 2.85, 0.42, 11, MUTED)
text(slide, "经营结果：获客依赖平台，服务却无法持续跟进。", 0.85, 4.25, 7.3, 0.38, 18, NAVY, True)
text(slide, "平台分发效率" , 1.05, 5.25, 2.0, 0.22, 12, MUTED, True)
rect(slide, 3.0, 5.22, 2.9, 0.32, GOLD, True)
text(slide, "自有沉淀能力", 6.35, 5.25, 2.0, 0.22, 12, MUTED, True)
rect(slide, 8.25, 5.22, 1.1, 0.32, BLUE, True)
text(slide, "需要补上“关系经营”这一层。", 0.85, 6.15, 5.7, 0.3, 15, BLUE, True)

# 5. Why own app
slide = new_slide(prs)
title(slide, "04 / WHY APP", "为什么需要建设自主 App 阵地", 5)
add_phone(slide, "02-learning.png", 0.85, 1.65, 2.35, 4.92)
for i, (a, b) in enumerate([
    ("稳定触达", "Push、站内消息和消息中心形成自有触达链路"),
    ("关系沉淀", "圈子、评论、问答和学习记录成为长期资产"),
    ("服务承载", "课程、直播、投研、会员和订单统一进入品牌阵地"),
    ("合规可控", "品牌、内容、权限和风险提示由客户统一管理"),
]):
    bullet(slide, a, b, 4.35, 1.45 + i * 1.15, 7.8, [BLUE, RGBColor(38,151,125), GOLD, RGBColor(190,78,91)][i])
text(slide, "自有 App 不是替代平台，而是承接核心用户与核心服务。", 4.35, 6.12, 7.0, 0.32, 15, NAVY, True)

# 6. Competitor routes
slide = new_slide(prs, PALE)
title(slide, "05 / MARKET LANDSCAPE", "金融行业竞品 App 的产品路线", 6)
text(slide, "竞品普遍从“内容获客”走向“工具 + 服务 + 社群”的复合经营。", 0.85, 1.4, 8.8, 0.35, 15, NAVY, True)
routes = [
    ("投教内容型", "客户类型 A", "课程体系 + 行情 + 社群", BLUE),
    ("投顾服务型", "客户类型 B", "直播 + 指标 + 会员服务", GOLD),
    ("数据工具型", "客户类型 C", "行情 + 选股 + 决策工具", RGBColor(38,151,125)),
    ("AI / 社区型", "客户类型 D", "AI 助手 + 讲师社区 + 订阅", RGBColor(190,78,91)),
]
for i, (a, b, c, color) in enumerate(routes):
    x = 0.85 + (i % 2) * 6.05
    y = 2.05 + (i // 2) * 1.65
    rect(slide, x, y, 5.35, 1.22, WHITE, True, LINE)
    rect(slide, x, y, 0.15, 1.22, color, True)
    text(slide, a, x + 0.3, y + 0.18, 1.45, 0.24, 14, NAVY, True)
    text(slide, b, x + 1.95, y + 0.18, 1.75, 0.24, 12, color, True)
    text(slide, c, x + 0.3, y + 0.62, 4.6, 0.25, 11, MUTED)
text(slide, "共同趋势：内容是入口，圈子和消息是留存，会员与工具是服务承接。", 0.85, 5.75, 10.3, 0.35, 16, BLUE, True)

# 7. Architecture
slide = new_slide(prs)
title(slide, "06 / PRODUCT BLUEPRINT", "自有 App 的整体产品架构", 7)
layers = [("内容层", "课程 · 直播 · 资讯 · 研报", BLUE), ("互动层", "圈子 · 话题 · 评论 · 问答", RGBColor(38,151,125)), ("服务层", "行情 · 投顾 · 会员 · 订单", GOLD), ("运营层", "Push · 标签 · 召回 · 数据", RGBColor(190,78,91))]
for i, (a, b, c) in enumerate(layers):
    y = 1.75 + i * 0.92
    rect(slide, 0.9, y, 6.45, 0.65, WHITE, True, LINE)
    rect(slide, 1.08, y + 0.13, 0.12, 0.39, c, True)
    text(slide, a, 1.45, y + 0.17, 1.15, 0.2, 13, NAVY, True)
    text(slide, b, 2.8, y + 0.17, 3.8, 0.2, 12, MUTED)
text(slide, "用户进入 App", 8.8, 1.95, 2.5, 0.28, 15, NAVY, True, PP_ALIGN.CENTER)
for i, (label, color) in enumerate([("内容吸引", NAVY), ("圈子建立关系", BLUE), ("消息持续触达", GOLD), ("会员与服务转化", RGBColor(38,151,125))]):
    y = 2.5 + i * 0.9
    rect(slide, 8.6, y, 2.9, 0.62, color, True)
    text(slide, label, 8.85, y + 0.19, 2.4, 0.2, 13, NAVY if color == GOLD else WHITE, True, PP_ALIGN.CENTER)
text(slide, "平台负责分发，自有 App 负责经营闭环。", 0.9, 6.05, 6.2, 0.3, 16, BLUE, True)

# 8. Circle core
slide = new_slide(prs, PALE)
title(slide, "07 / PRIVATE CIRCLE", "圈子：金融私域的核心承载", 8)
add_phone(slide, "05-circles.png", 0.78, 1.55, 2.25, 5.18)
add_phone(slide, "06-circle-detail.png", 3.48, 1.55, 2.25, 5.18)
text(slide, "圈子不是简单群聊，而是结构化、可持续运营的内容社区。", 6.9, 1.5, 5.3, 0.35, 16, NAVY, True)
for i, (a, b) in enumerate([("分层运营", "官方圈子、课程圈子、讲师圈子、VIP 圈子"), ("内容沉淀", "动态、回放、研报、问答和精选内容长期留存"), ("关系经营", "圈主置顶、评论互动、话题讨论与成员活跃"), ("服务连接", "课程、直播、会员和投顾服务都能回流到圈子")]):
    bullet(slide, a, b, 6.9, 2.15 + i * 0.92, 5.3, [BLUE, GOLD, RGBColor(38,151,125), RGBColor(190,78,91)][i])

# 9. Circle interactions
slide = new_slide(prs)
title(slide, "08 / CIRCLE OPERATION", "圈子内的内容与用户互动", 9)
for i, (a, b, c, icon) in enumerate([
    ("内容发布", "观点、盘前策略、课程笔记、直播回放", BLUE, "content"),
    ("互动反馈", "点赞、评论、回复、收藏和关注", RGBColor(38,151,125), "interaction"),
    ("主题运营", "话题、活动、打卡、问答和精选", GOLD, "topic"),
    ("关系转化", "圈友 → 学习用户 → 会员 / 投顾服务", RGBColor(190,78,91), "conversion"),
]):
    x = 0.85 + (i % 2) * 6.0
    y = 1.55 + (i // 2) * 1.7
    rect(slide, x, y, 5.25, 1.25, WHITE, True, LINE)
    slide.shapes.add_picture(str(make_icon(icon, c)), Inches(x + 0.14), Inches(y + 0.13), width=Inches(0.56), height=Inches(0.56))
    text(slide, a, x + 0.85, y + 0.2, 2.0, 0.25, 14, NAVY, True)
    text(slide, b, x + 0.85, y + 0.62, 3.95, 0.32, 11, MUTED)
text(slide, "圈子运营的目标：让用户有内容可看、有关系可互动、有服务可继续。", 0.85, 5.45, 10.5, 0.35, 16, BLUE, True)
text(slide, "示意页面：圈子详情承载内容、成员关系与服务入口。", 8.85, 5.65, 3.0, 0.5, 12, MUTED)

# 10. Push
slide = new_slide(prs, PALE)
title(slide, "09 / MESSAGE & PUSH", "消息推送：持续触达用户", 10)
add_phone(slide, "07-messages.png", 0.88, 1.55, 2.30, 5.18)
text(slide, "四类消息，覆盖用户从进入到复购的生命周期。", 4.25, 1.45, 7.5, 0.35, 16, NAVY, True)
messages = [("直播提醒", "预约、开播、回放更新", BLUE), ("圈子动态", "新观点、被回复、精选内容", RGBColor(38,151,125)), ("课程服务", "新课、续学、资料和学习进度", GOLD), ("运营召回", "沉默唤醒、会员到期、活动通知", RGBColor(190,78,91))]
for i, (a, b, c) in enumerate(messages):
    y = 2.18 + i * 0.9
    rect(slide, 4.25, y, 7.3, 0.62, WHITE, True, LINE)
    rect(slide, 4.45, y + 0.17, 0.28, 0.28, c, True)
    text(slide, a, 4.95, y + 0.16, 1.25, 0.2, 12, NAVY, True)
    text(slide, b, 6.55, y + 0.16, 4.55, 0.2, 11, MUTED)
text(slide, "平台触达不稳定 → 自有 App 通过 Push + 站内消息 + 用户标签建立可控链路。", 4.25, 6.1, 7.7, 0.3, 14, BLUE, True)

# 11. Segmentation
slide = new_slide(prs)
title(slide, "10 / USER OPERATIONS", "用户分层与召回运营", 11)
segments = [("新用户", "欢迎、首课、风险提示", BLUE), ("学习用户", "续学提醒、作业、训练营", GOLD), ("活跃圈友", "话题、直播、专家互动", RGBColor(38,151,125)), ("会员用户", "专属圈子、策略会、客服", RGBColor(190,78,91)), ("沉默用户", "内容召回、权益到期、活动", RGBColor(123,151,210))]
for i, (a, b, c) in enumerate(segments):
    x = 0.75 + i * 2.48
    rect(slide, x, 1.85, 2.05, 1.38, WHITE, True, LINE)
    rect(slide, x, 1.85, 2.05, 0.12, c, True)
    text(slide, a, x + 0.15, 2.2, 1.7, 0.24, 14, NAVY, True, PP_ALIGN.CENTER)
    text(slide, b, x + 0.17, 2.65, 1.7, 0.35, 10, MUTED, False, PP_ALIGN.CENTER)
text(slide, "标签来源", 0.85, 4.05, 1.0, 0.22, 12, NAVY, True)
text(slide, "学习进度 / 圈子行为 / 直播预约 / 会员状态 / 消息点击", 2.0, 4.05, 8.0, 0.22, 12, MUTED)
text(slide, "运营动作", 0.85, 4.75, 1.0, 0.22, 12, NAVY, True)
text(slide, "内容推荐 → 圈子互动 → Push 触达 → 服务转化 → 复购召回", 2.0, 4.75, 8.0, 0.22, 12, BLUE, True)
text(slide, "从“群发所有人”升级为“对的人，在合适的时机，收到合适的内容”。", 0.85, 6.05, 9.8, 0.3, 16, NAVY, True)

# 12. Linkage
slide = new_slide(prs, PALE)
title(slide, "11 / CONTENT TO SERVICE", "内容、直播、圈子和会员联动", 12)
for i, (a, b, c) in enumerate([("内容", "课程 / 资讯 / 研报", BLUE), ("直播", "预约 / 互动 / 回放", GOLD), ("圈子", "讨论 / 问答 / 关系", RGBColor(38,151,125)), ("会员", "权益 / 工具 / 服务", RGBColor(190,78,91))]):
    x = 0.85 + i * 3.05
    rect(slide, x, 2.2, 2.45, 1.15, c, True)
    text(slide, a, x, 2.48, 2.45, 0.25, 17, NAVY if c == GOLD else WHITE, True, PP_ALIGN.CENTER)
    text(slide, b, x, 2.9, 2.45, 0.22, 11, NAVY if c == GOLD else WHITE, False, PP_ALIGN.CENTER)
    if i < 3:
        text(slide, "→", x + 2.58, 2.6, 0.3, 0.25, 18, MUTED, True, PP_ALIGN.CENTER)
add_phone(slide, "03-course-detail.png", 1.0, 4.55, 1.2, 2.55)
add_phone(slide, "04-course-player.png", 2.8, 4.55, 1.2, 2.55)
add_phone(slide, "08-market.png", 4.6, 4.55, 1.2, 2.55)
add_phone(slide, "09-profile.png", 6.4, 4.55, 1.2, 2.55)
text(slide, "Demo 示意：学习、圈子、消息与会员在同一产品内形成连续服务。", 8.35, 5.1, 3.5, 0.6, 13, NAVY, True)

# 13. AI capability
slide = new_slide(prs)
title(slide, "12 / AI SERVICE", "AI 助手：接入客户 Agent，或定制金融服务 Agent", 13)
text(slide, "把客户已经拥有的内容、知识库和服务流程，转成 App 内可用的服务入口。", 0.85, 1.42, 10.5, 0.35, 15, NAVY, True)
for i, (heading, copy, color) in enumerate([
    ("接入客户现有 Agent", "复用客户知识库、Agent 编排与权限体系；App 提供品牌化入口、用户身份和服务上下文。", BLUE),
    ("定制金融服务 Agent", "围绕课程、直播回放、圈子内容、FAQ 和服务 SOP 定制内容检索、学习辅助与服务分流能力。", RGBColor(38,151,125)),
]):
    x = 0.85 + i * 5.85
    rect(slide, x, 2.08, 5.25, 1.48, WHITE, True, LINE)
    rect(slide, x, 2.08, 5.25, 0.13, color, True)
    text(slide, heading, x + 0.25, 2.38, 4.55, 0.25, 15, NAVY, True)
    text(slide, copy, x + 0.25, 2.8, 4.55, 0.5, 10, MUTED)
text(slide, "Demo 中的服务体验", 0.85, 4.18, 3.0, 0.25, 15, NAVY, True)
for i, (heading, copy, color) in enumerate([
    ("内容检索与总结", "定位课程、直播、圈子和已授权知识内容", BLUE),
    ("课程与直播辅助", "梳理重点、关联资料、继续学习与任务引导", GOLD),
    ("服务引导与人工转接", "按客户 SOP 分流至客服、顾问或企微服务", RGBColor(38,151,125)),
]):
    x = 0.85 + i * 3.75
    rect(slide, x, 4.65, 3.3, 1.15, WHITE, True, LINE)
    rect(slide, x + 0.2, 4.9, 0.36, 0.36, color, True)
    text(slide, heading, x + 0.72, 4.85, 2.25, 0.22, 12, NAVY, True)
    text(slide, copy, x + 0.2, 5.38, 2.8, 0.28, 9, MUTED)
text(slide, "AI 用于信息整理、学习辅助与服务引导；不构成投资建议或收益承诺。", 0.85, 6.32, 9.8, 0.28, 12, RGBColor(143, 92, 41), True)

# 14. Brand / backend / custom development / compliance
slide = new_slide(prs)
title(slide, "13 / PLATFORM CAPABILITY", "品牌、后台与合规能力", 14)
for i, (a, b, c) in enumerate([
    ("品牌能力", "Logo、主题色、首页 Banner、Tab 和页面风格可定制", BLUE),
    ("后台能力", "内容、课程、圈子、直播、消息和用户标签统一运营", GOLD),
    ("定制开发", "根据客户业务流程，提供页面、接口、数据和运营能力的定制开发", RGBColor(38,151,125)),
    ("合规能力", "隐私协议、风险提示、执业信息、权限和内容审核可配置", RGBColor(190,78,91)),
]):
    y = 1.45 + i * 1.08
    rect(slide, 0.85, y, 5.3, 0.98, WHITE, True, LINE)
    rect(slide, 1.08, y + 0.24, 0.14, 0.5, c, True)
    text(slide, a, 1.55, y + 0.2, 1.4, 0.23, 14, NAVY, True)
    text(slide, b, 3.05, y + 0.2, 2.7, 0.42, 10, MUTED)
rect(slide, 7.0, 1.55, 4.55, 4.5, NAVY, True)
text(slide, "可配置的品牌 App", 7.65, 2.0, 3.3, 0.3, 19, WHITE, True, PP_ALIGN.CENTER)
text(slide, "品牌配置", 7.65, 2.85, 1.4, 0.22, 13, GOLD, True)
text(slide, "内容与模块配置", 9.3, 2.85, 1.7, 0.22, 13, RGBColor(180,196,224), True)
text(slide, "↓", 9.15, 3.3, 0.3, 0.25, 18, WHITE, True, PP_ALIGN.CENTER)
text(slide, "多端发布 / H5 与原生组合", 7.65, 3.8, 3.3, 0.25, 15, WHITE, True, PP_ALIGN.CENTER)
text(slide, "↓", 9.15, 4.35, 0.3, 0.25, 18, WHITE, True, PP_ALIGN.CENTER)
text(slide, "合规、数据与运营可控", 7.65, 4.85, 3.3, 0.25, 15, GOLD, True, PP_ALIGN.CENTER)
rect(slide, 0.85, 6.05, 5.3, 0.55, WHITE, True, LINE)
text(slide, "展示主题", 1.08, 6.2, 0.9, 0.18, 11, NAVY, True)
text(slide, "深海蓝  ·  暖金棕  ·  青绿科技", 2.05, 6.2, 2.35, 0.18, 10, MUTED)
rect(slide, 4.55, 6.15, 1.25, 0.3, BLUE, True)
text(slide, "关怀版（大字）", 4.64, 6.21, 1.08, 0.14, 8, WHITE, True, PP_ALIGN.CENTER)

# 15. Compliance, security and deployment
slide = new_slide(prs, PALE)
title(slide, "14 / COMPLIANCE & SECURITY", "合规展示、信息安全与部分私有化部署", 15)
text(slide, "把合规要求做成产品能力，把数据边界做成部署选项。", 0.85, 1.35, 8.5, 0.3, 15, NAVY, True)
security_items = [
    ("合规组件", "R3/C3 适当性文案、风险提示、执业信息、隐私协议与 SDK 披露", BLUE),
    ("安全控制", "传输与存储加密、最小权限、租户隔离、敏感操作审计留痕", RGBColor(38,151,125)),
    ("数据边界", "用户、业务数据和审计日志可按客户安全域部署与管理", GOLD),
    ("部署方式", "支持部分私有化部署；按客户安全边界组合部署，不等同于全量私有化", RGBColor(190,78,91)),
]
for i, (a, b, c) in enumerate(security_items):
    x = 0.85 + (i % 2) * 3.25
    y = 2.05 + (i // 2) * 1.35
    rect(slide, x, y, 2.8, 1.08, WHITE, True, LINE)
    rect(slide, x, y, 0.08, 1.08, c, True)
    text(slide, a, x + 0.25, y + 0.18, 2.25, 0.23, 14, NAVY, True)
    text(slide, b, x + 0.25, y + 0.52, 2.25, 0.4, 10, MUTED)

rect(slide, 7.45, 1.85, 4.25, 4.55, NAVY, True)
text(slide, "按安全边界组合部署", 7.9, 2.2, 3.35, 0.28, 18, WHITE, True, PP_ALIGN.CENTER)
rect(slide, 8.05, 2.85, 3.05, 0.72, RGBColor(27, 66, 137), True)
text(slide, "客户侧安全域", 8.28, 3.05, 2.6, 0.2, 13, WHITE, True, PP_ALIGN.CENTER)
text(slide, "用户 / 业务数据 / 审计日志", 8.28, 3.34, 2.6, 0.18, 9, RGBColor(190, 211, 240), False, PP_ALIGN.CENTER)
text(slide, "+", 9.38, 3.78, 0.25, 0.25, 17, GOLD, True, PP_ALIGN.CENTER)
rect(slide, 8.05, 4.22, 3.05, 0.72, RGBColor(30, 116, 111), True)
text(slide, "平台侧可选能力", 8.28, 4.42, 2.6, 0.2, 13, WHITE, True, PP_ALIGN.CENTER)
text(slide, "内容运营 / 消息 / 直播 / 版本服务", 8.28, 4.71, 2.6, 0.18, 9, RGBColor(195, 235, 226), False, PP_ALIGN.CENTER)
text(slide, "部署组合需结合客户主体、数据分类分级与法务/安全评审确认。", 7.78, 5.55, 3.55, 0.45, 10, RGBColor(190, 211, 240), False, PP_ALIGN.CENTER)

# 16. Implementation and outcomes
slide = new_slide(prs, NAVY)
text(slide, "实施路径与预期业务结果", 0.85, 0.8, 7.8, 0.5, 28, WHITE, True)
text(slide, "平台负责获客，自有 App 负责沉淀；平台负责分发，自有圈子负责经营。", 0.85, 1.55, 10.5, 0.32, 16, RGBColor(205,218,240))
phases = [("阶段 1", "品牌与基础阵地", "首页 / 学习 / 我的 / 登录 / 合规", BLUE), ("阶段 2", "私域运营能力", "圈子 / 消息 / Push / 直播预约", GOLD), ("阶段 3", "服务与定制开发", "投研 / 会员 / 标签 / 召回 / 定制需求开发", RGBColor(38,151,125))]
for i, (a, b, c, color) in enumerate(phases):
    x = 0.9 + i * 4.05
    rect(slide, x, 2.55, 3.35, 1.55, color, True)
    text(slide, a, x + 0.22, 2.82, 1.1, 0.22, 12, NAVY if color == GOLD else WHITE, True)
    text(slide, b, x + 0.22, 3.15, 2.85, 0.25, 15, NAVY if color == GOLD else WHITE, True)
    text(slide, c, x + 0.22, 3.58, 2.85, 0.3, 10, NAVY if color == GOLD else WHITE)
text(slide, "预期结果", 0.9, 5.0, 1.2, 0.24, 14, GOLD, True)
text(slide, "降低单一平台依赖 · 稳定核心用户触达 · 沉淀内容与关系资产 · 支撑长期会员与投顾服务", 2.25, 5.0, 9.5, 0.3, 13, WHITE, True)
text(slide, "金融客户需要保留一个稳定、可持续经营的核心用户阵地。", 0.9, 6.25, 9.0, 0.35, 18, WHITE, True)

prs.save(OUT)
print(OUT)
